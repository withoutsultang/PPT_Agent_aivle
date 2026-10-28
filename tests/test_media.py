"""Real FFmpeg smoke test, independent of OpenAI and SerpAPI."""

import json
import shutil
import struct
import subprocess
import wave
from pathlib import Path

import pytest
from PIL import Image

from utils import concat_videos_ffmpeg, ffprobe_duration, render_mp4


@pytest.mark.skipif(not shutil.which("ffmpeg") or not shutil.which("ffprobe"), reason="FFmpeg is required")
def test_real_encoding_and_concat_with_quoted_paths(tmp_path):
    directory = tmp_path / "lecture with ' quote"
    directory.mkdir()
    image = directory / "slide.png"
    Image.new("RGB", (321, 181), "white").save(image)
    audio = directory / "speech.wav"
    with wave.open(str(audio), "wb") as stream:
        stream.setnchannels(1)
        stream.setsampwidth(2)
        stream.setframerate(24000)
        stream.writeframes(struct.pack("<h", 0) * 7200)
    videos = []
    for i in range(2):
        output = directory / f"slide-{i}.mp4"
        render_mp4(str(image), str(audio), str(output), width=321, height=181)
        videos.append(str(output))
    final = directory / "final.mp4"
    concat_videos_ffmpeg(videos, str(final))
    assert 0.5 <= ffprobe_duration(str(final)) <= 1.2
    data = json.loads(subprocess.check_output([
        "ffprobe", "-v", "error", "-show_streams", "-of", "json", str(final),
    ]))
    video = next(s for s in data["streams"] if s["codec_type"] == "video")
    assert (video["width"], video["height"]) == (320, 180)
    assert any(s["codec_type"] == "audio" for s in data["streams"])
    listing = final.with_suffix(".ffconcat").read_text()
    assert listing.count("\nfile ") == 2
    assert "\\n" not in listing


@pytest.mark.skipif(not shutil.which("soffice") or not shutil.which("pdftoppm"), reason="LibreOffice and Poppler are required")
def test_real_ppt_rendering_reuses_pdf_between_slides(tmp_path, monkeypatch):
    from pptx import Presentation
    import utils
    ppt = Presentation()
    for i in range(2):
        slide = ppt.slides.add_slide(ppt.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i + 1}"
    source = tmp_path / "input.pptx"
    ppt.save(source)
    original_run = utils.subprocess.run
    conversions = []
    def run(command, **kwargs):
        if command[0] == "soffice":
            conversions.append(command)
        return original_run(command, **kwargs)
    monkeypatch.setattr(utils.subprocess, "run", run)
    base = {"pptx_path": str(source), "work_dir": str(tmp_path / "slides")}
    first = utils.export_slide_as_png({**base, "slide_index": 0}, dpi=72)
    second = utils.export_slide_as_png({**base, "slide_index": 1}, dpi=72)
    assert len(conversions) == 1
    assert Path(first["slide_image"]).is_file()
    assert Path(second["slide_image"]).is_file()
    assert first["slide_image"] != second["slide_image"]
