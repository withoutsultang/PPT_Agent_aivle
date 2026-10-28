"""Regression checks without remote APIs or paid generation."""

import json
import logging
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches

import agent_nodes as nodes
import pipeline
from prompts import build_script_messages
from utils import clean_text


def quiz_payload():
    return [{"question": f"질문 {i}?", "options": [f"{n}. 선택지 {n}" for n in range(1, 5)], "answer": "2. 선택지 2"} for i in range(6)]


@pytest.mark.parametrize("wrapper", [False, True])
def test_quiz_accepts_legacy_array_and_json_object(wrapper):
    quizzes = quiz_payload()
    raw = {"quizzes": quizzes} if wrapper else quizzes
    assert nodes.parse_quiz_set(json.dumps(raw)) == quizzes


@pytest.mark.parametrize("defect", ["count", "options", "answer", "duplicate", "empty", "numbering"])
def test_quiz_rejects_invalid_contract(defect):
    quizzes = quiz_payload()
    if defect == "count":
        quizzes.pop()
    elif defect == "options":
        quizzes[0]["options"].pop()
    elif defect == "answer":
        quizzes[0]["answer"] = "없는 보기"
    elif defect == "duplicate":
        quizzes[1]["question"] = quizzes[0]["question"]
    elif defect == "empty":
        quizzes[0]["question"] = "  "
    else:
        quizzes[0]["options"][0] = "잘못된 번호"
    with pytest.raises(ValueError):
        nodes.parse_quiz_set(json.dumps({"quizzes": quizzes}))


def test_single_slide_prompt_has_both_opening_and_closure():
    messages = build_script_messages({"titles": ["한 장"], "slide_index": 0})
    assert "한 장으로 구성" in messages[1]["content"]
    assert "마무리 끝인사" in messages[1]["content"]


def test_middle_prompt_interpolates_previous_context_and_next_title():
    messages = build_script_messages({"titles": ["A", "B", "NEXT_TOPIC"], "slide_index": 1, "all_scripts": ["PREVIOUS_SCRIPT"]})
    prompt = messages[1]["content"]
    assert "PREVIOUS_SCRIPT" in prompt
    assert "NEXT_TOPIC" in prompt
    assert "{previous_script" not in prompt


def test_last_slide_prompt_does_not_preview_next_slide():
    prompt = build_script_messages({"titles": ["A", "B"], "slide_index": 1, "all_scripts": ["before"]})[1]["content"]
    assert "다음 내용 예고는 금지" in prompt
    assert "마무리 끝인사" in prompt


@pytest.mark.parametrize("index", [-1, 2])
def test_invalid_script_index_is_rejected(index):
    with pytest.raises(ValueError):
        build_script_messages({"titles": ["A"], "slide_index": index})


def test_whitespace_is_normalized_not_literal_backslash_s():
    assert clean_text("  A\n\t B  ") == "A B"
    assert clean_text(r"path\source") == r"path\source"


@pytest.mark.parametrize("speed", [0, -1, float("nan"), float("inf"), 5])
def test_invalid_speed_is_rejected_before_media_processing(speed):
    with pytest.raises(ValueError):
        nodes.atempo_filter(speed)


def test_atempo_filter_supports_chains():
    assert nodes.atempo_filter(4) == "atempo=2.0,atempo=2.0"
    assert nodes.atempo_filter(0.25) == "atempo=0.5,atempo=0.5"


def test_search_filters_domains_and_keeps_sources_aligned(monkeypatch):
    monkeypatch.setenv("SERPAPI_API_KEY", "TEST_ONLY")
    payload = {"organic_results": [
        {"title": "skip", "link": "https://writer.tistory.com/one", "snippet": "x"},
        {"title": "doc", "link": "https://example.org/doc", "snippet": "Reference snippet"},
        {"title": "bad", "link": "javascript:alert(1)", "snippet": "x"},
    ]}
    monkeypatch.setattr(nodes.requests, "get", lambda *a, **kw: SimpleNamespace(raise_for_status=lambda: None, json=lambda: payload))
    state = nodes.node_tool_search({"slide_index": 0, "titles": ["topic"], "texts": ["text"]})
    external = state["external_content"]
    assert len(external["references"]) == 1
    assert external["references"][0]["title"] == external["summaries"][0]["source"]


def test_search_failure_never_logs_key_in_exception(monkeypatch, caplog):
    monkeypatch.setenv("SERPAPI_API_KEY", "TEST_ONLY")
    def fail(*args, **kwargs):
        raise nodes.requests.RequestException("https://example.org?api_key=DO_NOT_LOG_THIS")
    monkeypatch.setattr(nodes.requests, "get", fail)
    with caplog.at_level(logging.WARNING):
        assert nodes.serpapi_search_by_title("title") == []
    assert "DO_NOT_LOG_THIS" not in caplog.text


def test_missing_current_video_cannot_shift_into_next_slide(tmp_path):
    later = tmp_path / "later.mp4"
    later.write_bytes(b"placeholder")
    with pytest.raises(ValueError):
        nodes.node_accumulate_and_step({"slide_index": 1, "slide_videos": {2: str(later)}})


def test_import_ui_and_build_without_api_key(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    from app import build_demo, display_quizzes, display_answers
    assert build_demo() is not None
    assert "\n" in display_quizzes(quiz_payload())
    assert "\\n" not in display_answers(quiz_payload())


def make_pptx(path, count):
    ppt = Presentation()
    for i in range(count):
        slide = ppt.slides.add_slide(ppt.slide_layouts[5])
        slide.shapes.title.text = f"Topic {i + 1}"
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1))
        shape.text = f"Shape context {i + 1}"
    ppt.save(path)


@pytest.mark.parametrize("count", [1, 3, 20])
def test_actual_langgraph_pipeline_offline(tmp_path, monkeypatch, count):
    """Real PPT parser, state channels, prompts, accumulation, quiz and job runner.

    Replace only external LLM/TTS/rendering boundaries with deterministic fakes.
    """
    source = tmp_path / "lesson.pptx"
    make_pptx(source, count)
    monkeypatch.setenv("OPENAI_API_KEY", "OFFLINE_TEST_ONLY")
    monkeypatch.delenv("SERPAPI_API_KEY", raising=False)
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "jobs"))
    monkeypatch.setattr(pipeline, "require_media_tools", lambda: None)
    monkeypatch.setattr(pipeline, "ffprobe_duration", lambda p: 1.0)
    requests = []
    def completion(**kwargs):
        requests.append(kwargs)
        if "response_format" in kwargs:
            content = json.dumps({"quizzes": quiz_payload()})
        elif isinstance(kwargs["messages"][1]["content"], list):
            content = "설명문입니다."
        else:
            content = "연속 강의 대본입니다."
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])
    monkeypatch.setattr(nodes, "get_client", lambda: SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=completion))))
    def snapshot(state):
        target = Path(state["work_dir"]) / f"{state['slide_index']}.png"
        target.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (64, 64)).save(target)
        return {**state, "slide_image": str(target)}
    monkeypatch.setattr(nodes, "export_slide_as_png", snapshot)
    def tts(state):
        target = Path(state["work_dir"]) / f"{state['slide_index']}.wav"
        target.write_bytes(b"offline audio")
        return {**state, "audio": str(target)}
    monkeypatch.setattr(pipeline, "node_tts", tts)
    def render(image, audio, output):
        Path(output).parent.mkdir(parents=True, exist_ok=True)
        Path(output).write_bytes(b"offline video")
    monkeypatch.setattr(nodes, "render_mp4", render)
    def concat(paths, output, **kwargs):
        assert len(paths) == count
        assert [Path(p).name for p in paths] == [f"slide-{i + 1:03d}.mp4" for i in range(count)]
        Path(output).write_bytes(b"offline final")
    monkeypatch.setattr(nodes, "concat_videos_ffmpeg", concat)
    result = pipeline.run_pipeline(str(source))
    assert result["slide_index"] == result["total_slides"] == count
    assert len(result["all_scripts"]) == len(result["video_paths"]) == count
    assert len(result["quiz_set"]) == 6
    assert "Shape context 1" in result["shape_texts"][0]
    assert Path(result["work_dir"], "scripts", "slide-001.txt").exists()
    assert Path(result["work_dir"], "result.json").exists()
    # Retry in the same second: workspaces must still be different.
    again = pipeline.run_pipeline(SimpleNamespace(name=str(source)))
    assert result["work_dir"] != again["work_dir"]


def test_quiz_failure_keeps_completed_video(tmp_path, monkeypatch):
    monkeypatch.setattr(nodes, "get_client", lambda: SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw: SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"quizzes":[]}'))])))))
    result = nodes.node_generate_quiz({"all_scripts": ["대본"], "work_dir": str(tmp_path), "final_video": "completed.mp4"})
    assert result["final_video"] == "completed.mp4"
    assert result["quiz_set"] == []
    assert result["notices"]


@pytest.mark.parametrize("error", [ValueError("PPTX를 선택해 주세요."), RuntimeError("PROVIDER_PRIVATE_DATA")])
def test_ui_failure_resets_outputs_without_leaking_provider_data(monkeypatch, error):
    import app
    def fail(*args, **kwargs):
        raise error
    monkeypatch.setattr(app, "run_pipeline", fail)
    result = app.generate_state_and_run(None, "tone", "alloy", "style", 60, 1)
    assert result[0] is None
    assert result[1]["visible"] is False
    assert result[4] == []
    assert "PROVIDER_PRIVATE_DATA" not in result[5]
